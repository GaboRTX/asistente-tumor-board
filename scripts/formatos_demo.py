# -*- coding: utf-8 -*-
"""
Distribucion de los casos de demostracion en formatos heterogeneos.

Los documentos de un caso real no llegan en un unico formato: el informe
anatomopatologico se emite en PDF, la historia clinica en Word, el perfil
molecular en una planilla. Distribuir los casos de demostracion de esa misma
manera permite ejercitar la ingesta tal como se la usara.

Las figuras incorporadas son esquemas rotulados como sinteticos. No se
generan imagenes que imiten estudios radiologicos reales: una figura con
apariencia de resonancia podria desprenderse de su contexto y ser tomada por
un estudio autentico.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Dict, List, Tuple


# Formato asignado a cada documento de cada caso. La variedad es deliberada:
# el conjunto debe ejercitar todos los lectores.
DISTRIBUCION: Dict[str, Dict[str, str]] = {
    "CASO-001": {"historia_clinica": "docx", "anatomia_patologica": "pdf",
                 "genomica": "xlsx", "imagenes": "txt"},
    "CASO-002": {"historia_clinica": "odt", "anatomia_patologica": "docx",
                 "genomica": "csv", "imagenes": "pdf"},
    "CASO-003": {"historia_clinica": "rtf", "anatomia_patologica": "txt",
                 "genomica": "md", "imagenes": "html"},
    "CASO-004": {"historia_clinica": "html", "anatomia_patologica": "docx",
                 "genomica": "xlsx", "imagenes": "pdf"},
    "CASO-005": {"historia_clinica": "txt", "anatomia_patologica": "txt",
                 "genomica": "txt", "imagenes": "txt"},
    "CASO-006": {"historia_clinica": "docx", "anatomia_patologica": "pdf",
                 "genomica": "xlsx", "imagenes": "md"},
    "CASO-007": {"historia_clinica": "pdf", "anatomia_patologica": "docx",
                 "genomica": "csv", "imagenes": "pdf"},
    "CASO-008": {"historia_clinica": "odt", "anatomia_patologica": "txt",
                 "genomica": "md", "imagenes": "docx"},
}

# Casos cuyo documento en Word incorpora figuras.
CON_FIGURAS = {"CASO-001", "CASO-004", "CASO-006", "CASO-008"}

LEYENDA = ("Documento sintetico de demostracion. No corresponde a un "
           "paciente real.")


# ---------------------------------------------------------------------------
# Figuras esquematicas
# ---------------------------------------------------------------------------

def figura_esquematica(destino: Path, titulo: str, pie: str) -> Path:
    """
    Genera un esquema rotulado como sintetico.

    Deliberadamente abstracto: formas geometricas y texto, sin apariencia de
    estudio radiologico. Una figura verosimil podria desprenderse de su
    contexto y ser tomada por un estudio autentico.
    """
    from PIL import Image, ImageDraw

    ancho, alto = 900, 620
    imagen = Image.new("RGB", (ancho, alto), "white")
    trazo = ImageDraw.Draw(imagen)

    trazo.rectangle([8, 8, ancho - 8, alto - 8], outline=(40, 40, 40), width=3)
    trazo.text((34, 34), "ESQUEMA SIN VALOR DIAGNOSTICO", fill=(140, 20, 20))
    trazo.text((34, 56), LEYENDA.upper(), fill=(140, 20, 20))
    trazo.text((34, 96), titulo, fill=(20, 20, 20))

    # Corte axial esquematico: dos elipses concéntricas y una region marcada.
    cx, cy = ancho // 2, 330
    trazo.ellipse([cx - 230, cy - 160, cx + 230, cy + 160],
                  outline=(40, 40, 40), width=3)
    trazo.ellipse([cx - 200, cy - 135, cx + 200, cy + 135],
                  outline=(120, 120, 120), width=1)
    trazo.line([cx, cy - 160, cx, cy + 160], fill=(120, 120, 120), width=1)

    # Lesion esquematica en el hemisferio izquierdo de la imagen.
    trazo.ellipse([cx - 160, cy - 80, cx - 40, cy + 20],
                  outline=(20, 20, 20), width=3)
    trazo.ellipse([cx - 140, cy - 62, cx - 62, cy + 2], fill=(200, 200, 200))
    trazo.text((cx - 155, cy + 34), "A: region marcada", fill=(20, 20, 20))

    trazo.text((34, 528), pie, fill=(20, 20, 20))
    trazo.text((34, 556), "Figura esquematica generada por programa. "
                          "No reproduce un estudio real.", fill=(90, 90, 90))
    trazo.text((34, 584), "Las mediciones constan en el texto del informe.",
               fill=(90, 90, 90))

    destino.parent.mkdir(parents=True, exist_ok=True)
    imagen.save(destino)
    return destino


# ---------------------------------------------------------------------------
# Escritores por formato
# ---------------------------------------------------------------------------

def _titulo_y_cuerpo(contenido: str) -> Tuple[str, List[str]]:
    lineas = [l.rstrip() for l in contenido.strip().splitlines()]
    titulo = lineas[0] if lineas else "Documento"
    return titulo, lineas[1:]


def escribir_txt(destino: Path, contenido: str, **_) -> Path:
    destino.write_text(contenido, encoding="utf-8")
    return destino


def escribir_md(destino: Path, contenido: str, **_) -> Path:
    titulo, cuerpo = _titulo_y_cuerpo(contenido)
    partes = [f"# {titulo}", ""]
    for linea in cuerpo:
        # Los encabezados en mayusculas pasan a encabezados de seccion.
        if linea and linea == linea.upper() and len(linea) > 3 and \
                not linea[0].isdigit():
            partes += ["", f"## {linea.title()}", ""]
        else:
            partes.append(linea)
    destino.write_text("\n".join(partes), encoding="utf-8")
    return destino


def escribir_html(destino: Path, contenido: str, **_) -> Path:
    import html as _html

    titulo, cuerpo = _titulo_y_cuerpo(contenido)
    parrafos = []
    for linea in cuerpo:
        if not linea.strip():
            continue
        if linea == linea.upper() and len(linea) > 3:
            parrafos.append(f"  <h2>{_html.escape(linea.title())}</h2>")
        else:
            parrafos.append(f"  <p>{_html.escape(linea)}</p>")

    destino.write_text(
        "<!doctype html>\n<html lang=\"es\">\n<head>\n"
        "  <meta charset=\"utf-8\">\n"
        f"  <title>{_html.escape(titulo)}</title>\n"
        "  <style>body{font-family:Arial,sans-serif;max-width:48rem;"
        "margin:2rem auto;line-height:1.5}</style>\n"
        "</head>\n<body>\n"
        f"  <h1>{_html.escape(titulo)}</h1>\n"
        + "\n".join(parrafos)
        + "\n</body>\n</html>\n", encoding="utf-8")
    return destino


def escribir_rtf(destino: Path, contenido: str, **_) -> Path:
    def escapar(texto: str) -> str:
        salida = []
        for caracter in texto:
            if caracter in "\\{}":
                salida.append("\\" + caracter)
            elif ord(caracter) > 127:
                salida.append(f"\\u{ord(caracter)}?")
            else:
                salida.append(caracter)
        return "".join(salida)

    cuerpo = "\\par\n".join(escapar(l) for l in contenido.strip().splitlines())
    destino.write_text(
        "{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}\\fs22\n"
        + cuerpo + "\n}", encoding="utf-8")
    return destino


def escribir_csv(destino: Path, contenido: str, **_) -> Path:
    """
    Vuelca el documento como planilla de dos columnas.

    Los informes de genomica se emiten con frecuencia en planilla, con una
    fila por determinacion. Las lineas con 'clave: valor' se separan en
    columnas; las demas se conservan en la primera.
    """
    with destino.open("w", encoding="utf-8", newline="") as f:
        escritor = csv.writer(f)
        escritor.writerow(["campo", "valor"])
        for linea in contenido.strip().splitlines():
            if ":" in linea and not linea.endswith(":"):
                clave, valor = linea.split(":", 1)
                escritor.writerow([clave.strip(), valor.strip()])
            elif linea.strip():
                escritor.writerow([linea.strip(), ""])
    return destino


def escribir_xlsx(destino: Path, contenido: str, **_) -> Path:
    import openpyxl

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Informe"
    hoja.append(["campo", "valor"])
    for linea in contenido.strip().splitlines():
        if ":" in linea and not linea.endswith(":"):
            clave, valor = linea.split(":", 1)
            hoja.append([clave.strip(), valor.strip()])
        elif linea.strip():
            hoja.append([linea.strip(), ""])
    hoja.column_dimensions["A"].width = 46
    hoja.column_dimensions["B"].width = 62
    libro.save(str(destino))
    return destino


def escribir_pdf(destino: Path, contenido: str, **_) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    ancho, alto = A4
    lienzo = canvas.Canvas(str(destino), pagesize=A4)
    margen, interlinea = 20 * mm, 12
    y = alto - margen

    util = ancho - 2 * margen

    def partir(texto: str, fuente: str, cuerpo: float) -> List[str]:
        """
        Divide la linea en renglones que entran en el ancho util.

        Truncarla destruiria informacion: una linea cortada puede perder
        justamente la negacion que invierte el sentido de un marcador.
        """
        if not texto:
            return [""]
        renglones, actual = [], ""
        for palabra in texto.split(" "):
            tentativa = f"{actual} {palabra}".strip()
            if lienzo.stringWidth(tentativa, fuente, cuerpo) <= util:
                actual = tentativa
            else:
                if actual:
                    renglones.append(actual)
                actual = palabra
        if actual:
            renglones.append(actual)
        return renglones

    for linea in contenido.strip().splitlines():
        es_titulo = linea == linea.upper() and len(linea.strip()) > 3
        fuente = "Helvetica-Bold" if es_titulo else "Helvetica"
        lienzo.setFont(fuente, 9.5)

        for renglon in partir(linea, fuente, 9.5):
            if y < margen:
                lienzo.showPage()
                lienzo.setFont(fuente, 9.5)
                y = alto - margen
            lienzo.drawString(margen, y, renglon)
            y -= interlinea

    lienzo.save()
    return destino


def escribir_odt(destino: Path, contenido: str, **_) -> Path:
    from odf.opendocument import OpenDocumentText
    from odf.text import H, P

    documento = OpenDocumentText()
    titulo, cuerpo = _titulo_y_cuerpo(contenido)
    documento.text.addElement(H(outlinelevel=1, text=titulo))
    for linea in cuerpo:
        documento.text.addElement(P(text=linea))
    documento.save(str(destino))
    return destino


def escribir_docx(destino: Path, contenido: str, figuras=None, **_) -> Path:
    """
    Compone un documento de Word con encabezados, una tabla y figuras.

    La tabla no es decorativa: los informes reales presentan el perfil
    molecular en forma tabular, y el sistema debe demostrar que lo incorpora
    al analisis.
    """
    import docx
    from docx.shared import Inches, Pt

    documento = docx.Document()
    estilo = documento.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(10)

    titulo, cuerpo = _titulo_y_cuerpo(contenido)
    documento.add_heading(titulo, level=1)

    filas_tabla: List[Tuple[str, str]] = []
    for linea in cuerpo:
        if not linea.strip():
            continue
        if linea == linea.upper() and len(linea.strip()) > 3:
            documento.add_heading(linea.title(), level=2)
        elif ":" in linea and len(linea) < 90 and not linea.endswith(":"):
            clave, valor = linea.split(":", 1)
            filas_tabla.append((clave.strip(), valor.strip()))
        else:
            documento.add_paragraph(linea)

    if filas_tabla:
        documento.add_heading("Resumen de determinaciones", level=2)
        tabla = documento.add_table(rows=len(filas_tabla) + 1, cols=2)
        tabla.style = "Table Grid"
        tabla.rows[0].cells[0].text = "Determinacion"
        tabla.rows[0].cells[1].text = "Resultado"
        for i, (clave, valor) in enumerate(filas_tabla, start=1):
            tabla.rows[i].cells[0].text = clave
            tabla.rows[i].cells[1].text = valor

    for figura in figuras or []:
        documento.add_paragraph()
        documento.add_picture(str(figura), width=Inches(5.2))
        documento.add_paragraph(
            "Figura esquematica sin valor diagnostico."
        ).italic = True

    documento.save(str(destino))
    return destino


ESCRITORES = {
    "txt": escribir_txt, "md": escribir_md, "html": escribir_html,
    "rtf": escribir_rtf, "csv": escribir_csv, "xlsx": escribir_xlsx,
    "pdf": escribir_pdf, "odt": escribir_odt, "docx": escribir_docx,
}


# ---------------------------------------------------------------------------
# Generacion
# ---------------------------------------------------------------------------

def generar(casos: Dict[str, Dict[str, str]], base: Path) -> Dict[str, List[str]]:
    """Escribe cada caso en los formatos que le asigna DISTRIBUCION."""
    resultado: Dict[str, List[str]] = {}

    for caso, documentos in sorted(casos.items()):
        directorio = base / caso
        directorio.mkdir(parents=True, exist_ok=True)
        formatos = DISTRIBUCION.get(caso, {})
        generados = []

        figuras: List[Path] = []
        if caso in CON_FIGURAS:
            carpeta = directorio / "figuras"
            figuras = [
                figura_esquematica(
                    carpeta / "esquema_lesion.png",
                    f"{caso} - esquema de localizacion",
                    "A: region de interes consignada en el informe."),
            ]

        for tipo, contenido in documentos.items():
            extension = formatos.get(tipo, "txt")
            escritor = ESCRITORES[extension]
            destino = directorio / f"{tipo}.{extension}"

            usa_figuras = figuras if (extension == "docx") else None
            escritor(destino, contenido, figuras=usa_figuras)
            generados.append(destino.name)

        resultado[caso] = sorted(generados)

    return resultado
