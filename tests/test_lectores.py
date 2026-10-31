# -*- coding: utf-8 -*-
"""
Pruebas de la lectura de documentos.

El riesgo central de esta capa no es la excepcion visible sino el silencio:
un documento que no pudo leerse y se omite sin aviso lleva al profesional a
suponer que el sistema lo considero. Varias pruebas verifican esa propiedad
de manera explicita.
"""

from pathlib import Path

import pytest

from core import lectores


TEXTO = ("INFORME ANATOMOPATOLOGICO\n"
         "Diagnostico integrado: Astrocitoma IDH mutante, grado 3 de la OMS.\n"
         "Promotor de MGMT: metilado.\n")


@pytest.fixture
def figura(tmp_path):
    from PIL import Image, ImageDraw

    ruta = tmp_path / "figura.png"
    imagen = Image.new("RGB", (420, 160), "white")
    trazo = ImageDraw.Draw(imagen)
    trazo.text((20, 40), "DOCUMENTO SINTETICO", fill="black")
    imagen.save(ruta)
    return ruta


# ---------------------------------------------------------------------------
# Formatos de texto
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("extension", [".txt", ".md", ".csv", ".json", ".log"])
def test_los_formatos_de_texto_se_leen(tmp_path, extension):
    ruta = tmp_path / f"doc{extension}"
    ruta.write_text(TEXTO, encoding="utf-8")
    assert "Astrocitoma" in lectores.leer(ruta).texto


def test_se_leen_codificaciones_distintas_de_utf8(tmp_path):
    """
    Los sistemas hospitalarios producen con frecuencia archivos en cp1252. Un
    fallo de codificacion no debe impedir la lectura.
    """
    ruta = tmp_path / "doc.txt"
    ruta.write_bytes("Resección quirúrgica completa.".encode("cp1252"))
    assert "resecci" in lectores.leer(ruta).texto.lower()


def test_el_html_descarta_el_marcado_y_los_scripts(tmp_path):
    ruta = tmp_path / "doc.html"
    ruta.write_text(
        "<html><body><h1>Informe</h1><p>IDH mutado.</p>"
        "<script>var x = 'no debe aparecer';</script></body></html>",
        encoding="utf-8")
    texto = lectores.leer(ruta).texto
    assert "IDH mutado" in texto
    assert "no debe aparecer" not in texto
    assert "<p>" not in texto


# ---------------------------------------------------------------------------
# Formatos de ofimatica
# ---------------------------------------------------------------------------

def test_el_docx_incorpora_el_texto_de_las_tablas(tmp_path):
    """
    Las tablas concentran el perfil molecular y los valores de laboratorio.
    Omitirlas perderia el dato de mayor densidad del documento.
    """
    import docx

    ruta = tmp_path / "doc.docx"
    documento = docx.Document()
    documento.add_paragraph("Informe de biologia molecular.")
    tabla = documento.add_table(rows=2, cols=2)
    tabla.rows[0].cells[0].text = "IDH"
    tabla.rows[0].cells[1].text = "mutado"
    tabla.rows[1].cells[0].text = "1p/19q"
    tabla.rows[1].cells[1].text = "codelecion ausente"
    documento.save(str(ruta))

    texto = lectores.leer(ruta).texto
    assert "IDH" in texto and "mutado" in texto
    assert "codelecion ausente" in texto


def test_el_docx_declara_las_imagenes_que_contiene(tmp_path, figura):
    """
    Una imagen embebida no se analiza. Declararla permite al profesional saber
    que ese contenido quedo fuera del analisis.
    """
    import docx

    ruta = tmp_path / "con_imagen.docx"
    documento = docx.Document()
    documento.add_paragraph("Estudio por imagenes.")
    documento.add_picture(str(figura))
    documento.save(str(ruta))

    lectura = lectores.leer(ruta)
    assert lectura.imagenes, "la imagen embebida debe declararse"
    assert any("no se analiza" in a for a in lectura.advertencias)


def test_la_planilla_se_lee_con_sus_hojas(tmp_path):
    import openpyxl

    ruta = tmp_path / "doc.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Marcadores"
    hoja.append(["marcador", "estado"])
    hoja.append(["TERT", "no mutado"])
    libro.save(str(ruta))

    texto = lectores.leer(ruta).texto
    assert "Marcadores" in texto
    assert "TERT" in texto and "no mutado" in texto


def test_el_odt_se_lee(tmp_path):
    from odf.opendocument import OpenDocumentText
    from odf.text import P

    ruta = tmp_path / "doc.odt"
    documento = OpenDocumentText()
    documento.text.addElement(P(text="Astrocitoma IDH mutante."))
    documento.save(str(ruta))

    assert "Astrocitoma" in lectores.leer(ruta).texto


def test_el_pdf_con_texto_se_lee(tmp_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    ruta = tmp_path / "doc.pdf"
    lienzo = canvas.Canvas(str(ruta), pagesize=A4)
    lienzo.drawString(60, 800, "Promotor de MGMT: metilado.")
    lienzo.save()

    lectura = lectores.leer(ruta)
    assert "MGMT" in lectura.texto
    assert lectura.paginas == 1


# ---------------------------------------------------------------------------
# Degradacion explicita
# ---------------------------------------------------------------------------

def test_sin_reconocimiento_optico_la_imagen_se_declara_sin_texto(
        tmp_path, figura, monkeypatch):
    """
    Esta es la propiedad central del modulo. Sin reconocimiento optico una
    imagen no aporta texto, y el sistema debe decirlo: presentarla como
    procesada induciria a suponer que su contenido fue considerado.
    """
    monkeypatch.setattr(lectores, "ocr_disponible", lambda: False)

    lectura = lectores.leer(figura)
    assert not lectura.tiene_texto
    assert lectura.imagenes == [figura.name]
    assert lectura.advertencias, "la imagen no leida debe advertirse"
    assert "no aporta texto" in lectura.advertencias[0]


def test_sin_reconocimiento_optico_el_pdf_escaneado_se_declara(tmp_path,
                                                               monkeypatch):
    """Un PDF sin texto incorporado es, casi siempre, un escaneo."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    monkeypatch.setattr(lectores, "ocr_disponible", lambda: False)

    ruta = tmp_path / "escaneado.pdf"
    lienzo = canvas.Canvas(str(ruta), pagesize=A4)
    lienzo.rect(100, 600, 200, 100)       # pagina sin texto
    lienzo.save()

    lectura = lectores.leer(ruta)
    assert not lectura.tiene_texto
    assert any("escaneado" in a.lower() for a in lectura.advertencias)


def test_un_archivo_ilegible_no_interrumpe_la_carga(tmp_path):
    """
    La carga de un caso con varios documentos no debe caerse por uno de ellos.
    El documento se registra sin texto y con el motivo.
    """
    ruta = tmp_path / "roto.docx"
    ruta.write_bytes(b"esto no es un documento de Word")

    lectura = lectores.leer(ruta)
    assert not lectura.tiene_texto
    assert lectura.advertencias


def test_el_formato_heredado_indica_la_conversion_necesaria(tmp_path):
    ruta = tmp_path / "antiguo.doc"
    ruta.write_bytes(b"\xd0\xcf\x11\xe0")

    lectura = lectores.leer(ruta)
    assert not lectura.tiene_texto
    assert ".docx" in lectura.advertencias[0]


def test_un_formato_ajeno_se_rechaza_nombrando_los_admitidos(tmp_path):
    ruta = tmp_path / "plano.dwg"
    ruta.write_text("x", encoding="utf-8")

    with pytest.raises(lectores.FormatoNoSoportado) as exc:
        lectores.leer(ruta)
    assert "pdf" in str(exc.value) and "docx" in str(exc.value)


def test_el_archivo_inexistente_se_distingue_del_ilegible(tmp_path):
    with pytest.raises(FileNotFoundError):
        lectores.leer(tmp_path / "no_existe.pdf")


# ---------------------------------------------------------------------------
# Integracion con la ingesta
# ---------------------------------------------------------------------------

def test_el_informe_de_lectura_expone_el_documento_sin_texto(tmp_path,
                                                             figura,
                                                             monkeypatch):
    """
    Un documento que no aporta fragmentos debe figurar igualmente en el
    informe. Omitirlo lo volveria invisible.
    """
    from core.ingesta import procesar_documentos_detallado

    monkeypatch.setattr(lectores, "ocr_disponible", lambda: False)

    patologia = tmp_path / "patologia.txt"
    patologia.write_text(TEXTO, encoding="utf-8")

    fragmentos, informes = procesar_documentos_detallado(
        "CASO-001", {"patologia": patologia, "imagenes": figura})

    assert len(informes) == 2
    imagen = next(i for i in informes if i["tipo_documento"] == "imagenes")
    assert imagen["fragmentos"] == 0
    assert not imagen["tiene_texto"]
    assert imagen["advertencias"]

    texto_doc = next(i for i in informes if i["tipo_documento"] == "patologia")
    assert texto_doc["fragmentos"] >= 1


# ---------------------------------------------------------------------------
# Composicion de filas tabulares
# ---------------------------------------------------------------------------

def test_la_fila_tabular_no_separa_la_determinacion_de_su_resultado():
    """
    Regresion grave. El analizador acota cada marcador a su clausula para no
    atribuirle el estado del vecino, y considera separadores de clausula a la
    coma, el punto y coma y la barra vertical.

    Unir la determinacion con su resultado mediante cualquiera de ellos los
    situaba en clausulas distintas. "EGFR | sin amplificacion" dejaba a EGFR
    sin estado y el marcador desaparecia del perfil.
    """
    from core.lectores import componer_fila
    from core.perfil_molecular import extraer_perfil

    linea = componer_fila(["EGFR", "sin amplificacion."])

    perfil = extraer_perfil([{"texto": linea, "tipo_documento": "genomica"}])
    assert "EGFR" in perfil.informados, f"linea compuesta: {linea!r}"
    assert perfil.informados["EGFR"].valor == "no amplificado"


def test_la_fila_de_tres_columnas_conserva_el_resultado():
    """Una fila con la tecnica intercalada no debe perder el resultado."""
    from core.lectores import componer_fila
    from core.perfil_molecular import extraer_perfil

    linea = componer_fila(["1p/19q", "FISH", "no se evidencia codelecion."])

    perfil = extraer_perfil([{"texto": linea, "tipo_documento": "genomica"}])
    assert perfil.informados["1p/19q"].valor == "codelecion ausente"


def test_la_coma_del_csv_no_se_confunde_con_puntuacion(tmp_path):
    """
    Regresion grave. La coma de un CSV es un separador estructural, no
    puntuacion de oracion. Entregarla como coma partia la linea
    "Codelecion 1p/19q,no se evidencia" y la negacion se perdia, de modo que
    un marcador ausente se consignaba como PRESENTE.

    Una inversion de polaridad es el peor desenlace posible: no se manifiesta
    como un error sino como un dato falso con apariencia de dato cierto.
    """
    from core.perfil_molecular import extraer_perfil

    ruta = tmp_path / "genomica.csv"
    ruta.write_text("campo,valor\n"
                    "Codelecion 1p/19q,no se evidencia.\n"
                    "Promotor de TERT,no se detecta mutacion.\n",
                    encoding="utf-8")

    texto = lectores.leer(ruta).texto
    perfil = extraer_perfil([{"texto": texto, "tipo_documento": "genomica"}])

    assert perfil.informados["1p/19q"].valor == "codelecion ausente", (
        "la negacion debe conservarse al cruzar el limite de columna")
    assert perfil.informados["TERT"].valor == "no mutado"


def test_la_celda_unica_se_conserva_tal_cual():
    from core.lectores import componer_fila

    assert componer_fila(["CONCLUSION"]) == "CONCLUSION"
    assert componer_fila([]) == ""
    assert componer_fila(["", "  "]) == ""
